"""Local artifact export helpers for agent tool calls.

The exporter intentionally keeps dependencies light. It uses python-docx when
available for DOCX and standard-library writers for Markdown, HTML, CSV, XLSX,
PPTX, and a simple text-first PDF.
"""

from __future__ import annotations

import csv
import html
import json
import re
import textwrap
import unicodedata
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape as xml_escape


EXPORT_DIR = Path(__file__).resolve().parent / "exports"
SUPPORTED_FORMATS = {"markdown", "md", "html", "docx", "pptx", "xlsx", "csv", "pdf"}


def export_artifact(params: dict[str, Any]) -> dict[str, Any]:
    """Export agent-provided content to a local artifact file."""

    title = str(params.get("title") or "Untitled Export").strip() or "Untitled Export"
    fmt = str(params.get("format") or "markdown").strip().lower()
    if fmt == "md":
        fmt = "markdown"
    if fmt not in SUPPORTED_FORMATS:
        return {
            "ok": False,
            "error": f"Unsupported format '{fmt}'. Supported formats: {sorted(SUPPORTED_FORMATS)}",
        }

    payload = _normalize_payload(params, title)
    suffix = "md" if fmt == "markdown" else fmt
    path = _unique_export_path(params.get("filename"), title, suffix)

    if fmt == "markdown":
        path.write_text(_to_markdown(payload), encoding="utf-8")
    elif fmt == "html":
        path.write_text(_to_html(payload), encoding="utf-8")
    elif fmt == "csv":
        _write_csv(path, payload)
    elif fmt == "xlsx":
        _write_xlsx(path, payload)
    elif fmt == "docx":
        _write_docx(path, payload)
    elif fmt == "pptx":
        _write_pptx(path, payload)
    elif fmt == "pdf":
        premium_result = _try_obsidian_pdf_export(params, payload)
        if premium_result:
            return premium_result
        _write_pdf(path, payload)

    return {
        "ok": True,
        "title": title,
        "format": suffix,
        "filename": path.name,
        "path": str(path),
        "download_url": f"/api/exports/{path.name}",
        "backend_url": f"/exports/{path.name}",
        "size_bytes": path.stat().st_size,
    }


def _try_obsidian_pdf_export(params: dict[str, Any], payload: dict[str, Any]) -> dict[str, Any] | None:
    try:
        from premium_export_tools import obsidian_export_pdf
    except Exception:
        return None

    markdown = str(params.get("content_markdown") or params.get("content") or "").strip()
    if not markdown:
        markdown = _to_markdown(payload)

    obsidian_params = dict(params)
    obsidian_params["title"] = payload["title"]
    obsidian_params["content_markdown"] = markdown
    result = obsidian_export_pdf(obsidian_params)
    if isinstance(result, dict) and result.get("ok"):
        result["format"] = "pdf"
        result["export_artifact_renderer"] = "obsidian-export"
        return result
    return None


def _normalize_payload(params: dict[str, Any], title: str) -> dict[str, Any]:
    content = str(params.get("content_markdown") or params.get("content") or "").strip()
    sections = _list_of_dicts(params.get("sections"))
    slides = _list_of_dicts(params.get("slides"))
    tables = _list_of_dicts(params.get("tables"))
    subtitle = str(params.get("subtitle") or params.get("audience") or "").strip()
    theme = str(params.get("theme") or "executive").strip().lower()

    if not sections and content:
        sections = [{"heading": "Content", "body": content}]
    if not slides:
        slides = _slides_from_sections(title, sections, content)
    if not tables:
        maybe_table = params.get("table")
        if isinstance(maybe_table, dict):
            tables = [maybe_table]

    return {
        "title": title,
        "subtitle": subtitle,
        "theme": theme if theme in {"consulting", "clinical", "technical", "executive"} else "executive",
        "content": content,
        "sections": sections,
        "slides": slides,
        "tables": [_normalize_table(t, idx) for idx, t in enumerate(tables, start=1)],
        "created_at": datetime.now(timezone.utc).isoformat(),
    }


def _list_of_dicts(value: Any) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        return []
    return [item for item in value if isinstance(item, dict)]


def _normalize_table(table: dict[str, Any], idx: int) -> dict[str, Any]:
    name = str(table.get("name") or table.get("title") or f"Table {idx}").strip()
    rows = table.get("rows")
    if not isinstance(rows, list):
        rows = []
    columns = table.get("columns")
    if not isinstance(columns, list) or not columns:
        columns = _infer_columns(rows)
    columns = [str(c) for c in columns]
    normalized_rows = []
    for row in rows:
        if isinstance(row, dict):
            normalized_rows.append([_cell(row.get(col, "")) for col in columns])
        elif isinstance(row, (list, tuple)):
            values = [_cell(v) for v in row]
            normalized_rows.append(values + [""] * max(0, len(columns) - len(values)))
        else:
            normalized_rows.append([_cell(row)])
    return {"name": name, "columns": columns, "rows": normalized_rows}


def _infer_columns(rows: list[Any]) -> list[str]:
    for row in rows:
        if isinstance(row, dict) and row:
            return [str(k) for k in row.keys()]
        if isinstance(row, (list, tuple)) and row:
            return [f"Column {i + 1}" for i in range(len(row))]
    return ["Content"]


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _slides_from_sections(title: str, sections: list[dict[str, Any]], content: str) -> list[dict[str, Any]]:
    if sections:
        overview = [
            _first_sentence(s.get("body", "") or " ".join(str(b) for b in s.get("bullets", []) if b))
            for s in sections[:4]
        ]
        slides = [{
            "title": title,
            "subtitle": "Presentation summary",
            "purpose": "Frame the document as a concise presentation.",
            "bullets": [b for b in overview if b][:4] or ["Summarizes the source document into a presentation narrative."],
            "takeaway": "Use the deck as a decision-ready summary, not as a report transcript.",
        }]
        summary_bullets = _summary_bullets_from_sections(sections, limit=5)
        if summary_bullets:
            slides.append({
                "title": "Executive Summary",
                "purpose": "Surface the main conclusions before the details.",
                "bullets": summary_bullets,
                "takeaway": summary_bullets[0],
            })
        for section in sections[:18]:
            body = str(section.get("body") or "")
            explicit_bullets = section.get("bullets")
            if isinstance(explicit_bullets, list) and explicit_bullets:
                bullets = [_fit_text(str(b), 105) for b in explicit_bullets[:4] if str(b).strip()]
            else:
                bullets = _summary_bullets_from_text(body, limit=4)
            slides.append({
                "title": _fit_text(str(section.get("heading") or "Section"), 58),
                "purpose": "Explain this section's role in the overall story.",
                "bullets": bullets,
                "takeaway": _fit_text(_first_sentence(body), 95),
                "notes": body[:1200],
            })
        return slides
    bullets = _summary_bullets_from_text(content)
    return [{
        "title": title,
        "subtitle": "Presentation summary",
        "purpose": "Summarize the provided document.",
        "bullets": bullets or ["No content provided."],
        "takeaway": bullets[0] if bullets else "",
    }]


def _first_sentence(text: Any) -> str:
    clean = _plain_text(str(text or ""))
    parts = re.split(r"(?<=[.!?])\s+", clean, maxsplit=1)
    return parts[0][:180] if parts and parts[0] else ""


def _bullets_from_text(text: str) -> list[str]:
    bullets = []
    for line in str(text or "").splitlines():
        stripped = line.strip().lstrip("-*0123456789. ")
        if stripped:
            bullets.append(stripped[:220])
        if len(bullets) >= 6:
            break
    if not bullets:
        wrapped = textwrap.wrap(_plain_text(text), width=110)
        bullets = [w[:220] for w in wrapped[:5]]
    return bullets


def _summary_bullets_from_sections(sections: list[dict[str, Any]], limit: int = 5) -> list[str]:
    bullets: list[str] = []
    for section in sections:
        heading = _plain_text(str(section.get("heading") or "")).strip()
        explicit = section.get("bullets")
        if isinstance(explicit, list) and explicit:
            seed = _plain_text(str(explicit[0]))
        else:
            seed = _first_sentence(section.get("body", ""))
        if heading and seed:
            bullet = f"{heading}: {seed}"
        else:
            bullet = seed or heading
        bullet = _fit_text(bullet, 155)
        if bullet and bullet not in bullets:
            bullets.append(bullet)
        if len(bullets) >= limit:
            break
    return bullets


def _summary_bullets_from_text(text: str, limit: int = 5) -> list[str]:
    candidates = _sentence_candidates(text)
    bullets: list[str] = []
    for sentence in candidates:
        clean = _fit_text(sentence, 155)
        if clean and clean not in bullets:
            bullets.append(clean)
        if len(bullets) >= limit:
            break
    if bullets:
        return bullets
    return _bullets_from_text(text)[:limit]


def _sentence_candidates(text: str) -> list[str]:
    clean = re.sub(r"\|.*\|", " ", str(text or ""))
    clean = re.sub(r"^#+\s+.*$", " ", clean, flags=re.MULTILINE)
    clean = re.sub(r"^\s*[-*]\s+", "", clean, flags=re.MULTILINE)
    clean = _plain_text(clean)
    parts = re.split(r"(?<=[.!?])\s+", clean)
    scored: list[tuple[int, int, str]] = []
    keywords = (
        "growth", "market", "clinical", "evidence", "regional", "vendor",
        "adoption", "risk", "gap", "recommend", "trend", "priority", "need",
    )
    for idx, part in enumerate(parts):
        sentence = part.strip()
        if len(sentence) < 35 or len(sentence) > 260:
            continue
        lower = sentence.lower()
        score = sum(1 for keyword in keywords if keyword in lower)
        if any(token in lower for token in ("http://", "https://", "doi:", "references")):
            score -= 2
        scored.append((-score, idx, sentence))
    scored.sort()
    return [sentence for _, _, sentence in scored]


def _plain_text(text: str) -> str:
    text = re.sub(r"`([^`]*)`", r"\1", text)
    text = re.sub(r"[*_#>\[\]]", "", text)
    return re.sub(r"\s+", " ", text).strip()


THEMES = {
    "executive": {
        "accent": "2456A6",
        "accent_dark": "173B75",
        "text": "17202A",
        "muted": "5B677A",
        "soft": "EEF3FA",
    },
    "consulting": {
        "accent": "1F7A6D",
        "accent_dark": "115149",
        "text": "18201F",
        "muted": "5F6F6B",
        "soft": "EEF7F5",
    },
    "clinical": {
        "accent": "0F7EA8",
        "accent_dark": "075D7D",
        "text": "13212A",
        "muted": "60717B",
        "soft": "EAF7FB",
    },
    "technical": {
        "accent": "5B5F97",
        "accent_dark": "363A6D",
        "text": "1B1D2A",
        "muted": "62677F",
        "soft": "F0F1FA",
    },
}


def _theme(payload: dict[str, Any]) -> dict[str, str]:
    return THEMES.get(str(payload.get("theme") or "executive"), THEMES["executive"])


def _safe_slug(value: str, default: str = "artifact") -> str:
    slug = re.sub(r"[^A-Za-z0-9._-]+", "-", value.strip()).strip("-._")
    return (slug or default)[:80]


def _unique_export_path(filename: Any, title: str, suffix: str) -> Path:
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    raw_name = str(filename or "").strip()
    if raw_name:
        stem = _safe_slug(Path(raw_name).stem)
    else:
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        stem = f"{_safe_slug(title)}-{stamp}"
    path = EXPORT_DIR / f"{stem}.{suffix}"
    counter = 2
    while path.exists():
        path = EXPORT_DIR / f"{stem}-{counter}.{suffix}"
        counter += 1
    return path


def _to_markdown(payload: dict[str, Any]) -> str:
    parts = [f"# {payload['title']}", ""]
    if payload.get("subtitle"):
        parts.extend([str(payload["subtitle"]), ""])
    if payload["content"]:
        parts.extend([payload["content"], ""])
    for section in payload["sections"]:
        heading = str(section.get("heading") or "Section").strip()
        body = str(section.get("body") or "").strip()
        parts.extend([f"## {heading}", "", body, ""])
    for table in payload["tables"]:
        parts.extend([f"## {table['name']}", "", _markdown_table(table), ""])
    return "\n".join(parts).strip() + "\n"


def _markdown_table(table: dict[str, Any]) -> str:
    columns = table["columns"]
    rows = table["rows"]
    header = "| " + " | ".join(_md_cell(c) for c in columns) + " |"
    sep = "| " + " | ".join("---" for _ in columns) + " |"
    body = ["| " + " | ".join(_md_cell(v) for v in row[:len(columns)]) + " |" for row in rows]
    return "\n".join([header, sep] + body)


def _md_cell(value: Any) -> str:
    return str(value).replace("|", "\\|").replace("\n", " ")


def _to_html(payload: dict[str, Any]) -> str:
    theme = _theme(payload)
    sections_html = []
    if payload.get("subtitle"):
        sections_html.append(f"<p class=\"subtitle\">{html.escape(str(payload['subtitle']))}</p>")
    if payload["content"]:
        sections_html.append(f"<section>{_markdown_to_html(payload['content'])}</section>")
    for section in payload["sections"]:
        sections_html.append(
            "<section>"
            f"<h2>{html.escape(str(section.get('heading') or 'Section'))}</h2>"
            f"{_markdown_to_html(str(section.get('body') or ''))}"
            "</section>"
        )
    for table in payload["tables"]:
        sections_html.append(f"<section><h2>{html.escape(table['name'])}</h2>{_html_table(table)}</section>")
    return f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{html.escape(payload['title'])}</title>
  <style>
    body {{ font-family: Arial, sans-serif; max-width: 920px; margin: 48px auto; padding: 0 28px; color: #{theme['text']}; line-height: 1.55; }}
    h1, h2 {{ line-height: 1.2; }}
    h1 {{ border-bottom: 2px solid #{theme['accent']}; padding-bottom: 12px; }}
    .subtitle {{ color: #{theme['muted']}; font-size: 18px; margin-top: -8px; }}
    table {{ border-collapse: collapse; width: 100%; margin: 18px 0; }}
    th, td {{ border: 1px solid #d8dee9; padding: 8px 10px; text-align: left; vertical-align: top; }}
    th {{ background: #{theme['soft']}; }}
    section {{ margin: 30px 0; }}
  </style>
</head>
<body>
  <h1>{html.escape(payload['title'])}</h1>
  {''.join(sections_html)}
</body>
</html>
"""


def _markdown_to_html(text: str) -> str:
    lines = text.splitlines()
    out = []
    in_list = False
    for line in lines:
        stripped = line.strip()
        if not stripped:
            if in_list:
                out.append("</ul>")
                in_list = False
            continue
        if stripped.startswith("#"):
            if in_list:
                out.append("</ul>")
                in_list = False
            level = min(len(stripped) - len(stripped.lstrip("#")), 4)
            out.append(f"<h{level}>{html.escape(stripped.lstrip('#').strip())}</h{level}>")
        elif stripped.startswith(("-", "*")):
            if not in_list:
                out.append("<ul>")
                in_list = True
            out.append(f"<li>{html.escape(stripped[1:].strip())}</li>")
        else:
            if in_list:
                out.append("</ul>")
                in_list = False
            out.append(f"<p>{html.escape(stripped)}</p>")
    if in_list:
        out.append("</ul>")
    return "\n".join(out)


def _html_table(table: dict[str, Any]) -> str:
    header = "".join(f"<th>{html.escape(c)}</th>" for c in table["columns"])
    rows = []
    for row in table["rows"]:
        rows.append("<tr>" + "".join(f"<td>{html.escape(str(v))}</td>" for v in row[:len(table["columns"])]) + "</tr>")
    return f"<table><thead><tr>{header}</tr></thead><tbody>{''.join(rows)}</tbody></table>"


def _write_csv(path: Path, payload: dict[str, Any]) -> None:
    table = payload["tables"][0] if payload["tables"] else {
        "columns": ["Content"],
        "rows": [[line] for line in _to_markdown(payload).splitlines() if line.strip()],
    }
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(table["columns"])
        writer.writerows(table["rows"])


def _write_docx(path: Path, payload: dict[str, Any]) -> None:
    try:
        from docx import Document
    except Exception as exc:  # pragma: no cover - depends on environment
        raise RuntimeError("DOCX export requires python-docx to be installed") from exc

    document = Document()
    document.add_heading(payload["title"], 0)
    if payload.get("subtitle"):
        document.add_paragraph(str(payload["subtitle"]))
    if payload["content"]:
        _add_markdown_to_docx(document, payload["content"])
    for section in payload["sections"]:
        document.add_heading(str(section.get("heading") or "Section"), level=1)
        _add_markdown_to_docx(document, str(section.get("body") or ""))
    for table in payload["tables"]:
        document.add_heading(table["name"], level=1)
        doc_table = document.add_table(rows=1, cols=max(1, len(table["columns"])))
        doc_table.style = "Table Grid"
        for idx, col in enumerate(table["columns"]):
            doc_table.rows[0].cells[idx].text = col
        for row in table["rows"]:
            cells = doc_table.add_row().cells
            for idx, value in enumerate(row[:len(cells)]):
                cells[idx].text = str(value)
    document.save(path)


def _add_markdown_to_docx(document: Any, text: str) -> None:
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("### "):
            document.add_heading(stripped[4:].strip(), level=3)
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:].strip(), level=2)
        elif stripped.startswith("# "):
            document.add_heading(stripped[2:].strip(), level=1)
        elif stripped.startswith(("- ", "* ")):
            document.add_paragraph(stripped[2:].strip(), style="List Bullet")
        else:
            document.add_paragraph(stripped)


def _write_xlsx(path: Path, payload: dict[str, Any]) -> None:
    tables = payload["tables"] or [{
        "name": "Content",
        "columns": ["Content"],
        "rows": [[line] for line in _to_markdown(payload).splitlines() if line.strip()],
    }]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", _xlsx_content_types(len(tables)))
        zf.writestr("_rels/.rels", _xlsx_root_rels())
        zf.writestr("xl/workbook.xml", _xlsx_workbook(tables))
        zf.writestr("xl/_rels/workbook.xml.rels", _xlsx_workbook_rels(len(tables)))
        zf.writestr("xl/styles.xml", _xlsx_styles())
        for idx, table in enumerate(tables, start=1):
            zf.writestr(f"xl/worksheets/sheet{idx}.xml", _xlsx_sheet(table))


def _xlsx_content_types(sheet_count: int) -> str:
    sheets = "\n".join(
        f'<Override PartName="/xl/worksheets/sheet{i}.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        for i in range(1, sheet_count + 1)
    )
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>
{sheets}
</Types>'''


def _xlsx_root_rels() -> str:
    return '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>'''


def _xlsx_workbook(tables: list[dict[str, Any]]) -> str:
    sheets = "\n".join(
        f'<sheet name="{xml_escape(_sheet_name(t["name"], idx))}" sheetId="{idx}" r:id="rId{idx}"/>'
        for idx, t in enumerate(tables, start=1)
    )
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<sheets>{sheets}</sheets>
</workbook>'''


def _xlsx_workbook_rels(sheet_count: int) -> str:
    rels = "\n".join(
        f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{i}.xml"/>'
        for i in range(1, sheet_count + 1)
    )
    rels += f'\n<Relationship Id="rId{sheet_count + 1}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{rels}</Relationships>'''


def _xlsx_styles() -> str:
    return '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>
<fills count="1"><fill><patternFill patternType="none"/></fill></fills>
<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>
<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>
<cellXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/></cellXfs>
</styleSheet>'''


def _xlsx_sheet(table: dict[str, Any]) -> str:
    rows = [table["columns"]] + table["rows"]
    row_xml = []
    for r_idx, row in enumerate(rows, start=1):
        cells = []
        for c_idx, value in enumerate(row, start=1):
            ref = f"{_column_name(c_idx)}{r_idx}"
            cells.append(f'<c r="{ref}" t="inlineStr"><is><t>{xml_escape(str(value))}</t></is></c>')
        row_xml.append(f'<row r="{r_idx}">{"".join(cells)}</row>')
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<sheetData>{"".join(row_xml)}</sheetData>
</worksheet>'''


def _sheet_name(name: str, idx: int) -> str:
    clean = re.sub(r"[\[\]:*?/\\]", " ", name).strip() or f"Sheet {idx}"
    return clean[:31]


def _column_name(index: int) -> str:
    name = ""
    while index:
        index, rem = divmod(index - 1, 26)
        name = chr(65 + rem) + name
    return name


def _write_pptx(path: Path, payload: dict[str, Any]) -> None:
    slides = payload["slides"] or _slides_from_sections(payload["title"], payload["sections"], payload["content"])
    slides = slides[:60]
    prepared_slides = []
    for idx, slide in enumerate(slides, start=1):
        prepared = dict(slide)
        prepared.setdefault("_deck_title", payload["title"])
        prepared.setdefault("_deck_subtitle", payload.get("subtitle") or "")
        prepared.setdefault("_theme", payload.get("theme") or "executive")
        prepared.setdefault("_slide_count", len(slides))
        prepared.setdefault("_slide_number", idx)
        prepared_slides.append(prepared)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", _pptx_content_types(len(slides)))
        zf.writestr("_rels/.rels", _pptx_root_rels())
        zf.writestr("docProps/core.xml", _core_props(payload["title"]))
        zf.writestr("docProps/app.xml", _app_props(len(slides)))
        zf.writestr("ppt/presentation.xml", _pptx_presentation(len(slides)))
        zf.writestr("ppt/_rels/presentation.xml.rels", _pptx_presentation_rels(len(slides)))
        for idx, slide in enumerate(prepared_slides, start=1):
            zf.writestr(f"ppt/slides/slide{idx}.xml", _pptx_slide(slide, idx))


def _pptx_content_types(slide_count: int) -> str:
    slide_overrides = "\n".join(
        f'<Override PartName="/ppt/slides/slide{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
        for i in range(1, slide_count + 1)
    )
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>
<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>
<Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>
{slide_overrides}
</Types>'''


def _pptx_root_rels() -> str:
    return '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>
<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/>
</Relationships>'''


def _core_props(title: str) -> str:
    now = datetime.now(timezone.utc).isoformat()
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<dc:title>{xml_escape(title)}</dc:title>
<dc:creator>Obsidian AI</dc:creator>
<cp:lastModifiedBy>Obsidian AI</cp:lastModifiedBy>
<dcterms:created xsi:type="dcterms:W3CDTF">{now}</dcterms:created>
<dcterms:modified xsi:type="dcterms:W3CDTF">{now}</dcterms:modified>
</cp:coreProperties>'''


def _app_props(slide_count: int) -> str:
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">
<Application>Obsidian AI</Application><Slides>{slide_count}</Slides>
</Properties>'''


def _pptx_presentation(slide_count: int) -> str:
    sld_ids = "\n".join(f'<p:sldId id="{255 + i}" r:id="rId{i}"/>' for i in range(1, slide_count + 1))
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:presentation xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
<p:sldIdLst>{sld_ids}</p:sldIdLst>
<p:sldSz cx="9144000" cy="5143500" type="screen16x9"/>
<p:notesSz cx="6858000" cy="9144000"/>
</p:presentation>'''


def _pptx_presentation_rels(slide_count: int) -> str:
    rels = "\n".join(
        f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide{i}.xml"/>'
        for i in range(1, slide_count + 1)
    )
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{rels}</Relationships>'''


def _pptx_slide(slide: dict[str, Any], idx: int) -> str:
    title = _fit_text(str(slide.get("title") or f"Slide {idx}"), 58)
    theme = THEMES.get(str(slide.get("_theme") or "executive"), THEMES["executive"])
    subtitle = _fit_text(str(slide.get("subtitle") or slide.get("purpose") or (slide.get("_deck_subtitle") if idx == 1 else "") or "").strip(), 92)
    takeaway = _fit_text(str(slide.get("takeaway") or "").strip(), 95)
    visual = str(slide.get("visual") or "").strip()
    bullets = slide.get("bullets")
    if not isinstance(bullets, list):
        bullets = _bullets_from_text(str(slide.get("body") or slide.get("notes") or ""))
    clean_bullets = [_fit_text(str(b).strip(), 105) for b in bullets[:4] if str(b).strip()]
    total_bullet_chars = sum(len(b) for b in clean_bullets)
    bullet_size = 1850 if total_bullet_chars < 280 else 1650
    title_size = 2750 if len(title) <= 42 else 2350
    bullet_xml = "".join(_pptx_paragraph(b, size=bullet_size, bullet=True) for b in clean_bullets)
    if not bullet_xml:
        bullet_xml = _pptx_paragraph("", size=bullet_size)
    subtitle_xml = _pptx_paragraph(subtitle, size=1350, color=theme["muted"]) if subtitle else ""
    takeaway_xml = _pptx_paragraph(takeaway, size=1400, bold=True, color=theme["accent_dark"]) if takeaway else ""
    visual_xml = _pptx_paragraph(_fit_text(f"Suggested visual: {visual}", 95), size=1250, color=theme["muted"]) if visual else ""
    footer = f"{idx}/{int(slide.get('_slide_count') or idx)}"
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
<p:cSld><p:spTree>
<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>
{_pptx_rect_shape(2, "Accent Bar", 0, 0, 152400, 5143500, theme["accent"])}
{_pptx_rect_shape(3, "Header Band", 152400, 0, 8991600, 1257300, theme["soft"])}
{_pptx_text_shape(4, "Title", 457200, 182880, 8229600, 487680, _pptx_paragraph(title, size=title_size, bold=True, color=theme["text"]), font_scale=72000)}
{_pptx_text_shape(5, "Subtitle", 502920, 777240, 7802880, 320040, subtitle_xml, font_scale=76000)}
{_pptx_text_shape(6, "Body", 731520, 1524000, 7589520, 2194560, bullet_xml, font_scale=70000)}
{_pptx_text_shape(7, "Takeaway", 731520, 3931920, 7589520, 426720, takeaway_xml or visual_xml, font_scale=76000)}
{_pptx_text_shape(8, "Footer", 7924800, 4754880, 914400, 274320, _pptx_paragraph(footer, size=1100, color=theme["muted"]))}
</p:spTree></p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr>
</p:sld>'''


def _fit_text(text: str, limit: int) -> str:
    clean = _plain_text(text)
    if len(clean) <= limit:
        return clean
    return clean[: max(0, limit - 1)].rstrip() + "..."


def _pptx_rect_shape(shape_id: int, name: str, x: int, y: int, cx: int, cy: int, fill: str) -> str:
    return f'''<p:sp>
<p:nvSpPr><p:cNvPr id="{shape_id}" name="{name}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="{fill}"/></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr>
</p:sp>'''


def _pptx_text_shape(shape_id: int, name: str, x: int, y: int, cx: int, cy: int, paragraphs: str, font_scale: int = 85000) -> str:
    if not paragraphs:
        paragraphs = _pptx_paragraph("", size=1200)
    autofit = f'<a:normAutofit fontScale="{font_scale}" lnSpcReduction="20000"/>'
    return f'''<p:sp>
<p:nvSpPr><p:cNvPr id="{shape_id}" name="{name}"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>
<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></p:spPr>
<p:txBody><a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0">{autofit}</a:bodyPr><a:lstStyle/>{paragraphs}</p:txBody>
</p:sp>'''


def _pptx_paragraph(text: str, size: int = 2200, bold: bool = False, color: str = "17202A", bullet: bool = False) -> str:
    bold_attr = ' b="1"' if bold else ""
    ppr = '<a:pPr marL="342900" indent="-171450"><a:buChar char="&#8226;"/></a:pPr>' if bullet else ""
    return f'<a:p>{ppr}<a:r><a:rPr lang="en-US" sz="{size}"{bold_attr}><a:solidFill><a:srgbClr val="{color}"/></a:solidFill></a:rPr><a:t>{xml_escape(text)}</a:t></a:r><a:endParaRPr lang="en-US" sz="{size}"/></a:p>'


def _write_pdf(path: Path, payload: dict[str, Any]) -> None:
    pages = _pdf_pages(payload)
    _write_simple_pdf(path, payload, pages)


def _pdf_pages(payload: dict[str, Any]) -> list[list[dict[str, Any]]]:
    story = _pdf_story(payload)
    pages: list[list[dict[str, Any]]] = [[]]
    y = 708
    bottom = 62

    for block in story:
        size = int(block["size"])
        leading = int(size * 1.45)
        gap = int(block.get("gap", 0))
        width = 42 if size >= 18 else 68 if size >= 12 else 88
        lines = textwrap.wrap(block["text"], width=width) or [""]
        required = len(lines) * leading + gap
        if pages[-1] and y - required < bottom:
            pages.append([])
            y = 708
        for line in lines:
            pages[-1].append({**block, "text": line, "y": y})
            y -= leading
        y -= gap
    return pages or [[]]


def _pdf_story(payload: dict[str, Any]) -> list[dict[str, Any]]:
    story: list[dict[str, Any]] = []

    def add(text: Any, size: int = 10, bold: bool = False, gap: int = 4) -> None:
        clean = _plain_text(str(text or "")).strip()
        if clean:
            story.append({"text": clean, "size": size, "bold": bold, "gap": gap})

    add(payload["title"], 20, True, 10)
    if payload.get("subtitle"):
        add(payload["subtitle"], 11, False, 16)

    for section in payload["sections"]:
        add(section.get("heading") or "Section", 14, True, 6)
        bullets = section.get("bullets")
        if isinstance(bullets, list) and bullets:
            for bullet in bullets:
                add(f"- {bullet}", 10, False, 2)
            continue
        for line in str(section.get("body") or "").splitlines():
            stripped = line.strip()
            if not stripped:
                continue
            if stripped.startswith("### "):
                add(stripped[4:], 11, True, 4)
            elif stripped.startswith("## "):
                add(stripped[3:], 12, True, 5)
            elif stripped.startswith("# "):
                add(stripped[2:], 13, True, 6)
            elif stripped.startswith(("- ", "* ")):
                add(f"- {stripped[2:]}", 10, False, 2)
            elif stripped.startswith("|") and stripped.endswith("|"):
                cells = [c.strip() for c in stripped.strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    add(" | ".join(cells), 9, False, 2)
            else:
                add(stripped, 10, False, 4)

    for table in payload["tables"]:
        add(table["name"], 13, True, 5)
        add(" | ".join(table["columns"]), 9, True, 2)
        for row in table["rows"][:80]:
            add(" | ".join(str(v) for v in row[:len(table["columns"])]), 9, False, 1)

    return story


def _write_simple_pdf(path: Path, payload: dict[str, Any], pages: list[list[dict[str, Any]]]) -> None:
    objects: list[bytes] = []

    def add(obj: str | bytes) -> int:
        objects.append(obj.encode("latin-1", errors="replace") if isinstance(obj, str) else obj)
        return len(objects)

    catalog_id = add("")
    pages_id = add("")
    font_id = add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    bold_font_id = add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>")
    page_ids = []
    for page_number, page_lines in enumerate(pages, start=1):
        stream = _pdf_page_stream(payload, page_lines, page_number, len(pages))
        content_id = add(f"<< /Length {len(stream)} >>\nstream\n".encode("latin-1") + stream + b"\nendstream")
        page_id = add(f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 {font_id} 0 R /F2 {bold_font_id} 0 R >> >> /Contents {content_id} 0 R >>")
        page_ids.append(page_id)
    objects[catalog_id - 1] = f"<< /Type /Catalog /Pages {pages_id} 0 R >>".encode("latin-1")
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects[pages_id - 1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode("latin-1")

    content = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for idx, obj in enumerate(objects, start=1):
        offsets.append(len(content))
        content.extend(f"{idx} 0 obj\n".encode("latin-1"))
        content.extend(obj)
        content.extend(b"\nendobj\n")
    xref_at = len(content)
    content.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("latin-1"))
    for offset in offsets[1:]:
        content.extend(f"{offset:010d} 00000 n \n".encode("latin-1"))
    content.extend(f"trailer\n<< /Size {len(objects) + 1} /Root {catalog_id} 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode("latin-1"))
    path.write_bytes(bytes(content))


def _pdf_page_stream(payload: dict[str, Any], lines: list[dict[str, Any]], page_number: int, page_count: int) -> bytes:
    theme = _theme(payload)
    accent = _pdf_rgb(theme["accent"])
    muted = _pdf_rgb(theme["muted"])
    header_title = _fit_text(_pdf_latin_text(payload["title"]), 74)
    commands = [
        "q",
        f"{accent} rg",
        "0 770 612 22 re f",
        "Q",
        f"{muted} rg",
        "BT /F1 8 Tf 72 748 Td",
        f"({_pdf_escape(header_title)}) Tj",
        "ET",
        "BT /F1 8 Tf 500 34 Td",
        f"({_pdf_escape(f'{page_number}/{page_count}')}) Tj",
        "ET",
        "0 0 0 rg",
    ]
    for item in lines:
        font = "/F2" if item.get("bold") else "/F1"
        size = int(item.get("size") or 10)
        x = 72
        y = int(item.get("y") or 700)
        commands.append(f"BT {font} {size} Tf {x} {y} Td ({_pdf_escape(item['text'])}) Tj ET")
    return "\n".join(commands).encode("latin-1", errors="replace")


def _pdf_rgb(hex_color: str) -> str:
    value = hex_color.strip().lstrip("#")
    try:
        r = int(value[0:2], 16) / 255
        g = int(value[2:4], 16) / 255
        b = int(value[4:6], 16) / 255
    except Exception:
        r, g, b = 0, 0, 0
    return f"{r:.3f} {g:.3f} {b:.3f}"


def _pdf_escape(text: str) -> str:
    safe = _pdf_latin_text(text)
    return safe.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


_PDF_TEXT_REPLACEMENTS = str.maketrans({
    "\u2010": "-",
    "\u2011": "-",
    "\u2012": "-",
    "\u2013": "-",
    "\u2014": "-",
    "\u2015": "-",
    "\u2212": "-",
    "\u2022": "*",
    "\u2190": "<-",
    "\u2192": "->",
    "\u2194": "<->",
    "\u21d2": "=>",
    "\u00d7": "x",
    "\u00b1": "+/-",
    "\u2264": "<=",
    "\u2265": ">=",
    "\u2248": "~",
    "\u2018": "'",
    "\u2019": "'",
    "\u201c": '"',
    "\u201d": '"',
    "\u2026": "...",
    "\u00b5": "u",
    "\u03bc": "u",
})


def _pdf_latin_text(text: Any) -> str:
    translated = str(text or "").translate(_PDF_TEXT_REPLACEMENTS)
    normalized = unicodedata.normalize("NFKD", translated)
    return normalized.encode("latin-1", errors="ignore").decode("latin-1")
